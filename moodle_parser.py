import os
import time
import requests
import undetected_chromedriver as uc
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from pptx import Presentation
from pypdf import PdfReader

# ==========================================
# CONFIGURATION (FILL WITH YOUR DATA)
# ==========================================
EMAIL = "YOUR_EMAIL@studenti.unime.it"
PASSWORD = "YOUR_PASSWORD"
TARGET_URL = "https://elearning.unime.it/course/view.php?id=395" # Replace with your specific course URL

OPENROUTER_API_KEY = "YOUR_OPENROUTER_API_KEY"
AI_MODEL = "stealth/space-bunny-alpha" 

TELEGRAM_TOKEN = "YOUR_TELEGRAM_BOT_TOKEN"
TELEGRAM_CHAT_ID = "YOUR_TELEGRAM_CHAT_ID"

# --- LOOP SETTINGS ---
CHECK_INTERVAL_MINUTES = 15  
# -----------------------

PROFILE_FOLDER = os.path.abspath("moodle_bot_profile")
DOWNLOAD_FOLDER = os.path.abspath("moodle_downloads")
MEMORY_FILE = os.path.abspath("processed_lectures.txt")

if not os.path.exists(DOWNLOAD_FOLDER):
    os.makedirs(DOWNLOAD_FOLDER)

# ==========================================
# MEMORY SYSTEM
# ==========================================
def load_processed_links():
    if not os.path.exists(MEMORY_FILE):
        return []
    with open(MEMORY_FILE, "r") as f:
        return [line.strip() for line in f.readlines()]

def save_processed_link(url):
    with open(MEMORY_FILE, "a") as f:
        f.write(f"{url}\n")

# ==========================================
# CORE LOGIC
# ==========================================
def get_lecture_links():
    print(f"Loading profile from: {PROFILE_FOLDER}")
    options = uc.ChromeOptions()
    options.add_argument(f"--user-data-dir={PROFILE_FOLDER}")
    
    prefs = {
        "download.default_directory": DOWNLOAD_FOLDER,
        "download.prompt_for_download": False,
        "plugins.always_open_pdf_externally": True
    }
    options.add_experimental_option("prefs", prefs)
    
    driver = uc.Chrome(options=options) 
    wait = WebDriverWait(driver, 15)
    short_wait = WebDriverWait(driver, 5)

    lecture_links = []

    try:
        print("Navigating to Moodle...")
        driver.get(TARGET_URL)

        sso_button = wait.until(EC.element_to_be_clickable((By.CSS_SELECTOR, "a.login-identityprovider-btn")))
        driver.execute_script("arguments[0].click();", sso_button)
        time.sleep(3)

        try:
            email_field = short_wait.until(EC.element_to_be_clickable((By.NAME, "loginfmt")))
            email_field.clear()
            email_field.send_keys(EMAIL)
            driver.find_element(By.ID, "idSIButton9").click()
            time.sleep(2) 
            password_field = wait.until(EC.element_to_be_clickable((By.NAME, "passwd")))
            password_field.clear()
            password_field.send_keys(PASSWORD)
            time.sleep(1)
            driver.find_element(By.ID, "idSIButton9").click()
            time.sleep(2)
            try:
                wait.until(EC.element_to_be_clickable((By.ID, "idSIButton9"))).click()
            except: pass
        except Exception:
            print("Cookies active! Skipped login.")

        time.sleep(5) 
        all_links = driver.find_elements(By.TAG_NAME, "a")
        
        for link in all_links:
            href = link.get_attribute('href')
            text = link.text.strip()
            if href and "resource/view.php" in href:
                lecture_links.append({"title": text, "url": href, "element": link})

        return lecture_links, driver

    except Exception as e:
        print(f"Error in parsing: {e}")
        if driver: driver.quit()
        return [], None

def extract_text_from_file(filepath):
    print(f"Extracting text from: {filepath}")
    text_content = ""
    
    try:
        if filepath.lower().endswith('.pdf'):
            reader = PdfReader(filepath)
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text_content += extracted + "\n"
                    
        elif filepath.lower().endswith('.pptx'):
            prs = Presentation(filepath)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text_content += shape.text + "\n"
        else:
            print("Unsupported file format.")
            
        return text_content.strip()
    except Exception as e:
        print(f"Failed to read file: {e}")
        return ""

def ask_openrouter(text):
    print(f"Sending text to OpenRouter AI (Model: {AI_MODEL})...")
    url = "https://openrouter.ai/api/v1/chat/completions"
    headers = {
        "Authorization": f"Bearer {OPENROUTER_API_KEY}",
        "Content-Type": "application/json"
    }
    
    prompt = (
        "You are a strict academic assistant. Summarize the following university lecture into exactly 3 key points.\n"
        "Strict rules:\n"
        "1. Output MUST be in English.\n"
        "2. Use plain text ONLY. Do NOT use markdown, asterisks (**), or bold text.\n"
        "3. Start your response directly with the number '1.'. Do NOT output any conversational filler.\n\n"
        f"Lecture text:\n{text[:4000]}"
    )
    
    data = {
        "model": AI_MODEL,
        "messages": [{"role": "user", "content": prompt}]
    }

    try:
        response = requests.post(url, headers=headers, json=data)
        response.raise_for_status()
        return response.json()["choices"][0]["message"]["content"]
    except requests.exceptions.RequestException as e:
        error_details = e.response.text if e.response else str(e)
        print(f"AI API Error Details: {error_details}")
        return f"Error generating summary. API responded with: {e}"

def send_to_telegram(title, summary):
    url = f"https://api.telegram.org/bot{TELEGRAM_TOKEN}/sendMessage"
    message = f"📚 New Lecture: {title}\n\nAI Summary:\n{summary}"
    try:
        response = requests.post(url, data={"chat_id": TELEGRAM_CHAT_ID, "text": message})
        response.raise_for_status() 
        print("Successfully sent to Telegram.")
    except Exception as e:
        error_details = response.text if 'response' in locals() else str(e)
        print(f"Telegram API Error Details: {error_details}")

# ==========================================
# MAIN CONTINUOUS PIPELINE
# ==========================================
if __name__ == "__main__":
    print("--- STARTING MOODLE CONTINUOUS MONITOR ---")
    
    while True:
        current_time = time.strftime('%Y-%m-%d %H:%M:%S')
        print(f"\n==============================================")
        print(f"[CYCLE START] Checking for new lectures at {current_time}")
        print(f"==============================================\n")
        
        try:
            processed_links = load_processed_links()
            lectures, driver = get_lecture_links()
            
            if lectures and driver:
                latest_lec = lectures[-1]
                
                if latest_lec['url'] in processed_links:
                    print(f"Latest lecture '{latest_lec['title']}' is already processed. Everything is up to date.")
                else:
                    print(f"\n[NEW] Processing latest lecture: {latest_lec['title']}")
                    
                    for old_file in os.listdir(DOWNLOAD_FOLDER):
                        try: os.remove(os.path.join(DOWNLOAD_FOLDER, old_file))
                        except: pass

                    bypass_url = latest_lec['url']
                    bypass_url += "&redirect=1" if "?" in bypass_url else "?redirect=1"
                    
                    print("Opening lecture page...")
                    driver.get(bypass_url)
                    time.sleep(5) 
                    
                    print("Hunting for Amanote popups...")
                    clicked_amanote = False
                    
                    try:
                        btn = driver.find_element(By.XPATH, "//*[contains(normalize-space(text()), 'View here')]")
                        driver.execute_script("arguments[0].click();", btn)
                        print("[!] Clicked 'View here' on main page.")
                        clicked_amanote = True
                    except: pass

                    if not clicked_amanote:
                        for frame in driver.find_elements(By.TAG_NAME, "iframe"):
                            driver.switch_to.frame(frame)
                            try:
                                btn = driver.find_element(By.XPATH, "//*[contains(normalize-space(text()), 'View here')]")
                                driver.execute_script("arguments[0].click();", btn)
                                print("[!] Clicked 'View here' inside iframe! Amanote defeated.")
                                clicked_amanote = True
                            except: pass
                            finally:
                                driver.switch_to.default_content() 
                            if clicked_amanote: break
                    
                    if clicked_amanote:
                        time.sleep(4) 
                        
                    try:
                        real_file = driver.find_elements(By.XPATH, "//a[contains(@href, 'pluginfile.php')]")
                        if real_file:
                            driver.execute_script("arguments[0].click();", real_file[0])
                            print("[!] Found native Moodle file link. Clicked.")
                    except: pass

                    print("Waiting 15 seconds for download...")
                    time.sleep(15)
                    
                    files = [f for f in os.listdir(DOWNLOAD_FOLDER) if not f.endswith('.crdownload')]
                    
                    if files:
                        filepath = os.path.join(DOWNLOAD_FOLDER, files[0])
                        raw_text = extract_text_from_file(filepath)
                        
                        if raw_text:
                            summary = ask_openrouter(raw_text)
                            send_to_telegram(latest_lec['title'], summary)
                            save_processed_link(latest_lec['url'])
                        else:
                            print("No text found in file. Skipping memory save.")
                    else:
                        print("Download failed. File not found in folder.")
                        
                driver.quit()
            else:
                print("Initialization failed or no lectures found this cycle.")
                if driver: driver.quit()
                
        except Exception as e:
            print(f"\n[CRITICAL ERROR] Something went wrong during the cycle: {e}")
            try:
                if 'driver' in locals() and driver:
                    driver.quit()
            except: pass
            
        print(f"\nCycle finished. Sleeping for {CHECK_INTERVAL_MINUTES} minutes...")
        time.sleep(CHECK_INTERVAL_MINUTES * 60)